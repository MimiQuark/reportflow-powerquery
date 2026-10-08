from __future__ import annotations

import json
import sys
import unittest
from pathlib import Path

from openpyxl import load_workbook

BASE_DIR = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(BASE_DIR / "src"))

from validate_data import validate_data  # noqa: E402


class ExcelReportPipelineTest(unittest.TestCase):
    def test_source_data_quality(self):
        report = validate_data()
        self.assertTrue(report["passed"])

    def test_workbook_structure(self):
        path = BASE_DIR / "output" / "自动化经营报表-含PowerQuery.xlsx"
        self.assertTrue(path.exists(), "Run run.ps1 first")
        workbook = load_workbook(path, read_only=False, data_only=False)
        for sheet in ["项目说明", "经营看板", "PowerQuery说明", "字段口径", "数据异常", "数据透视汇总", "版本记录", "raw_sales", "raw_products", "raw_returns", "raw_targets", "fact_sales"]:
            self.assertIn(sheet, workbook.sheetnames)
        dashboard = workbook["经营看板"]
        self.assertGreaterEqual(len(workbook['数据透视汇总']._pivots), 1)
        self.assertGreaterEqual(len(dashboard._charts), 3)
        self.assertTrue(any(str(cell.value).startswith("=SUM(fact_sales[") for row in dashboard.iter_rows() for cell in row if cell.value))

    def test_power_query_embed_report(self):
        path = BASE_DIR / "output" / "powerquery_embed_report.json"
        self.assertTrue(path.exists(), "Run embed_power_queries.ps1 first")
        report = json.loads(path.read_text(encoding="utf-8-sig"))
        self.assertEqual(report["queryCount"], 6)
        self.assertEqual(set(report["queries"]), {"qry_fact_sales", "qry_monthly_kpi", "qry_category_kpi", "qry_region_kpi", "qry_return_summary", "qry_quality_issues"})

    def test_power_query_refresh_report(self):
        path = BASE_DIR / "output" / "powerquery_refresh_report.json"
        self.assertTrue(path.exists(), "Run test_power_queries.ps1 first")
        report = json.loads(path.read_text(encoding="utf-8-sig"))
        self.assertTrue(report["passed"])
        rows = {item["query"]: item["rows"] for item in report["results"]}
        self.assertGreater(rows["qry_fact_sales"], 40000)
        self.assertGreaterEqual(rows["qry_monthly_kpi"], 36)


if __name__ == "__main__":
    unittest.main()